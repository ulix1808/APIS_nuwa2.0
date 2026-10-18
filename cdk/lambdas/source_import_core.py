"""
Source file → records, for the background import worker (no database, no AWS: unit-testable).

Mirrors what the admin page did in the browser (app/admin/page.tsx in the BFF):
- read CSV / TSV (text kept as published: no date guessing), XLSX (date cells as dd/mm/yyyy) or JSON;
- header row chosen by the user; empty headers → columna_N, repeated → name_2;
- repair reversible text damage (Windows-1252 opened as Mac Roman, UTF-8 read as Latin-1), the same
  rules as lib/source-file-validation.ts;
- columns with the same role are joined with a space; accents stripped except Ñ; JSON key = the role
  ("id", "nombre", "pais", "fecha", "estado") or the custom label ("custom:Monto" → "Monto").

Each record also gets:
- row_hash: hash of the normalized content (dates ISO, amounts as numbers, no accents, case or empty
  fields), so a re-upload in another format is "unchanged";
- record_key: the normalized value of the key field (RFC, CURP, folio…), or "h:<row_hash>" without one.
"""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import json
import re
import unicodedata
from dataclasses import dataclass, field
from typing import Any, Iterable

from chunk_normalize import CHUNK_TEXT_NORMALIZED_KEY, prepare_chunk_text_for_storage

BUILTIN_ROLES = ("id", "nombre", "pais", "fecha", "estado")
JSON_WRAPPER_KEYS = ("data", "records", "items", "results", "rows", "entries", "list", "sources", "registros")


class SourceFileError(ValueError):
    """The file cannot be read as configured (message shown to the user)."""


# --- decoding --------------------------------------------------------------------------------------


def _looks_like_utf16le_ascii(b: bytes) -> bool:
    if len(b) < 12 or len(b) % 2:
        return False
    pairs = min(len(b) // 2, 256)
    ascii_like = 0
    for i in range(0, pairs * 2, 2):
        lo, hi = b[i], b[i + 1]
        if hi != 0:
            return False
        if 0x20 <= lo <= 0x7E or lo in (0x09, 0x0A, 0x0D):
            ascii_like += 1
    return ascii_like / pairs > 0.88


def decode_text_bytes(b: bytes) -> str:
    """UTF-8 (BOM or not), UTF-16, else Windows-1252 — as lib/spreadsheet-text-decode.ts."""
    if b.startswith(b"\xef\xbb\xbf"):
        return unicodedata.normalize("NFC", b[3:].decode("utf-8", errors="replace"))
    if b.startswith(b"\xff\xfe"):
        return unicodedata.normalize("NFC", b[2:].decode("utf-16-le", errors="replace"))
    if b.startswith(b"\xfe\xff"):
        return unicodedata.normalize("NFC", b[2:].decode("utf-16-be", errors="replace"))
    if _looks_like_utf16le_ascii(b):
        return unicodedata.normalize("NFC", b.decode("utf-16-le", errors="replace"))
    try:
        return unicodedata.normalize("NFC", b.decode("utf-8"))
    except UnicodeDecodeError:
        return unicodedata.normalize("NFC", b.decode("cp1252", errors="replace"))


# --- reading ---------------------------------------------------------------------------------------


def _cell_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        if (v.hour, v.minute, v.second) == (0, 0, 0):
            return v.strftime("%d/%m/%Y")
        return v.strftime("%d/%m/%Y %H:%M")
    if isinstance(v, dt.date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    return str(v).strip()


def read_matrix(data: bytes, file_name: str, *, sheet_name: str | None = None) -> list[list[str]]:
    """All rows of the file as text cells (CSV/TSV/TXT, XLSX, JSON handled by read_json_records)."""
    lower = file_name.lower()
    if lower.endswith(".xlsx") or lower.endswith(".xlsm"):
        return _read_xlsx(data, sheet_name)
    if lower.endswith((".csv", ".tsv", ".txt")):
        return _read_delimited(decode_text_bytes(data), tab=lower.endswith(".tsv"))
    raise SourceFileError(f"Formato no soportado para procesamiento en segundo plano: {file_name}")


def _read_delimited(text: str, *, tab: bool) -> list[list[str]]:
    if tab:
        delim = "\t"
    else:
        sample = text[:65536]
        try:
            delim = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            delim = ","
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text, newline=""), delimiter=delim)]


def _read_xlsx(data: bytes, sheet_name: str | None) -> list[list[str]]:
    import openpyxl  # bundled in cdk/lambdas

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise SourceFileError(f"No se pudo leer el Excel: {e}") from e
    try:
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]
        return [[_cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def read_json_records(data: bytes) -> tuple[list[str], list[dict[str, str]]]:
    """Array of objects or {data|records|items|…: [...]}; as lib/json-upload-parse.ts."""
    try:
        parsed = json.loads(decode_text_bytes(data))
    except json.JSONDecodeError as e:
        raise SourceFileError("JSON inválido. Use un arreglo de objetos o un objeto con clave data/records/items.") from e
    if isinstance(parsed, list):
        objects = [o for o in parsed if isinstance(o, dict)]
    elif isinstance(parsed, dict):
        objects = next((parsed[k] for k in JSON_WRAPPER_KEYS if isinstance(parsed.get(k), list)), [parsed])
        objects = [o for o in objects if isinstance(o, dict)]
    else:
        raise SourceFileError("JSON inválido.")
    headers: list[str] = []
    seen: set[str] = set()
    for o in objects:
        for k in o:
            k2 = str(k).strip()
            if k2 and k2 not in seen:
                seen.add(k2)
                headers.append(k2)

    def cell(v: Any) -> str:
        if v is None:
            return ""
        if isinstance(v, str):
            return v.strip()
        if isinstance(v, bool):
            return "true" if v else "false"
        if isinstance(v, (int, float)):
            return _cell_text(v)
        return json.dumps(v, ensure_ascii=False)

    rows = [{h: cell(o.get(h)) for h in headers} for o in objects]
    return headers, [r for r in rows if any(r.values())]


def table_from_matrix(matrix: list[list[str]], header_row_index: int) -> tuple[list[str], list[dict[str, str]]]:
    if not matrix:
        raise SourceFileError("El archivo está vacío.")
    idx = min(max(0, header_row_index), len(matrix) - 1)
    seen: dict[str, int] = {}
    headers: list[str] = []
    for i, h in enumerate(matrix[idx]):
        base = (h or "").strip() or f"columna_{i + 1}"
        n = seen.get(base, 0)
        seen[base] = n + 1
        headers.append(base if n == 0 else f"{base}_{n + 1}")
    rows = []
    for raw in matrix[idx + 1 :]:
        row = {h: (raw[i].strip() if i < len(raw) and raw[i] is not None else "") for i, h in enumerate(headers)}
        if any(row.values()):
            rows.append(row)
    return headers, rows


# --- text repair (same rules as lib/source-file-validation.ts) -------------------------------------

_MAC_MISREAD = {"—": "Ñ", "”": "Ó", "Û": "ó", "¡": "Á", "Ò": "ñ", "Ì": "í", "Õ": "Í", "…": "É", "„": "Ñ", "†": " "}
_MAC_MARKS = "".join(_MAC_MISREAD)
_LETTER = "A-Za-zÁÉÍÓÚáéíóúÑñ"
_MAC_BETWEEN_LETTERS = re.compile(f"[{_LETTER}][{re.escape(_MAC_MARKS)}][{_LETTER}]")
_UTF8_AS_LATIN1 = re.compile("[ÂÃ][\u0080-¿‘-›Œ-Ÿˆ˜€™ŠŽ]")


def detect_garbled(cells: Iterable[str]) -> str | None:
    mac = utf8 = 0
    for c in cells:
        if _MAC_BETWEEN_LETTERS.search(c):
            mac += 1
        if _UTF8_AS_LATIN1.search(c):
            utf8 += 1
        if mac >= 2 or utf8 >= 2:
            break
    if utf8 >= 2 and utf8 >= mac:
        return "utf8"
    if mac >= 2:
        return "mac"
    return None


def repair_text(s: str, damage: str | None) -> str:
    if not damage or not s:
        return s
    if damage == "mac":
        return "".join(_MAC_MISREAD.get(c, c) for c in s)
    try:
        return s.encode("cp1252").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return s


def strip_accents_keep_enye(s: str) -> str:
    s = s.replace("Ñ", "\0").replace("ñ", "\1")
    s = "".join(c for c in unicodedata.normalize("NFD", s) if not ("̀" <= c <= "ͯ"))
    return s.replace("\0", "Ñ").replace("\1", "ñ")


def header_match_key(h: str) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFD", h).encode("ascii", "ignore").decode()).strip().lower()


# --- normalization and hashing ---------------------------------------------------------------------

_MONTHS = {
    "ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6, "jul": 7, "ago": 8, "sep": 9, "sept": 9,
    "oct": 10, "nov": 11, "dic": 12, "jan": 1, "apr": 4, "aug": 8, "dec": 12,
}


def _fold(s: str) -> str:
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s).strip().lower()


def _iso_date(s: str) -> str | None:
    t = s.strip()
    m = re.fullmatch(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", t)
    if m:
        d, mo, y = map(int, m.groups())
    else:
        m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T].*)?", t)
        if m:
            y, mo, d = map(int, m.groups())
        else:
            m = re.fullmatch(r"(\d{1,2})[-/ ]([A-Za-z]{3,4})\.?[-/ ](\d{2,4})", t)
            if not m or m.group(2).lower() not in _MONTHS:
                return None
            d, mo, y = int(m.group(1)), _MONTHS[m.group(2).lower()], int(m.group(3))
            if y < 100:
                y += 2000
    try:
        return dt.date(y, mo, d).isoformat()
    except ValueError:
        return None


def _number(s: str) -> str | None:
    t = s.replace("$", "").replace(",", "").replace(" ", "")
    if not re.fullmatch(r"-?\d+(\.\d+)?", t) or len(t) > 18:
        return None
    if re.fullmatch(r"0\d+", t):
        return None  # folio / clave with leading zeros: keep as text
    v = float(t)
    return str(int(v)) if v.is_integer() else repr(v)


def normalize_value(v: str) -> str:
    return _iso_date(v) or _number(v) or _fold(v)


def row_hash(obj: dict[str, Any]) -> str:
    items = sorted(
        (_fold(k), normalize_value(str(v)))
        for k, v in obj.items()
        if k != CHUNK_TEXT_NORMALIZED_KEY and str(v).strip()
    )
    return hashlib.sha256(json.dumps(items, ensure_ascii=False).encode("utf-8")).hexdigest()[:32]


def normalize_key(v: str) -> str:
    return re.sub(r"[^0-9A-Z&Ñ]", "", strip_accents_keep_enye(str(v)).upper())


def record_key(obj: dict[str, Any], key_field: str | None, digest: str) -> str:
    if key_field:
        k = normalize_key(str(obj.get(key_field) or ""))
        if k:
            return k
    return f"h:{digest}"


# --- records ---------------------------------------------------------------------------------------


@dataclass
class Record:
    key: str
    digest: str
    chunk_text: str
    normalized: str


@dataclass
class BuildResult:
    records: list[Record]
    rows_in_file: int
    duplicates_in_file: int
    empty_rows: int
    repaired_cells: int
    damage: str | None
    lost_character_cells: int
    unmatched_mapping: list[str] = field(default_factory=list)


def json_key_for_role(role: str) -> str:
    if role.startswith("custom:"):
        return role[len("custom:") :].strip() or "custom_field"
    return role


def build_records(
    headers: list[str],
    rows: list[dict[str, str]],
    column_mapping: dict[str, str],
    key_field: str | None,
) -> BuildResult:
    """Rows → stored records. Same key twice in the file: the last row wins."""
    damage = detect_garbled([*headers, *(v for r in rows[:5000] for v in r.values() if v)])
    repaired = 0

    def fix(s: str) -> str:
        nonlocal repaired
        out = repair_text(s, damage)
        if out != s:
            repaired += 1
        return out

    fixed_headers = [fix(h) for h in headers]
    # The browser mapped the (possibly repaired) headers of a sample: match by folded name.
    by_key = {header_match_key(h): role for h, role in column_mapping.items()}
    roles = [by_key.get(header_match_key(fh)) or by_key.get(header_match_key(h)) for h, fh in zip(headers, fixed_headers)]
    mapped = {header_match_key(fh) for fh, r in zip(fixed_headers, roles) if r}
    unmatched = [h for h, r in column_mapping.items() if r and r != "ignorar" and header_match_key(h) not in mapped]

    plan: list[tuple[str, list[str]]] = []  # (json key, source headers) in first-column order
    seen_keys: dict[str, int] = {}
    for h, role in zip(headers, roles):
        if not role or role == "ignorar":
            continue
        jk = json_key_for_role(role)
        if jk not in seen_keys:
            seen_keys[jk] = len(plan)
            plan.append((jk, []))
        plan[seen_keys[jk]][1].append(h)

    out: dict[str, Record] = {}
    dups = empty = lost = 0
    for row in rows:
        obj: dict[str, str] = {}
        for jk, src in plan:
            parts = [fix(row.get(h, "")).strip() for h in src]
            obj[jk] = strip_accents_keep_enye(" ".join(p for p in parts if p))
        if not any(obj.values()):
            empty += 1
            continue
        if any("�" in v for v in obj.values()):
            lost += 1
        stored, norm = prepare_chunk_text_for_storage(json.dumps(obj, ensure_ascii=False))
        digest = row_hash(obj)
        key = record_key(obj, key_field, digest)
        if key in out:
            dups += 1
            del out[key]  # keep file order of the last occurrence
        out[key] = Record(key, digest, stored, norm)
    return BuildResult(
        records=list(out.values()),
        rows_in_file=len(rows),
        duplicates_in_file=dups,
        empty_rows=empty,
        repaired_cells=repaired,
        damage=damage,
        lost_character_cells=lost,
        unmatched_mapping=unmatched,
    )


def key_and_hash_of_stored(chunk_text: str, key_field: str | None) -> tuple[str, str]:
    """record_key / row_hash for a row already in the table (loaded before keys existed)."""
    try:
        obj = json.loads(chunk_text)
        if not isinstance(obj, dict):
            raise ValueError
    except (json.JSONDecodeError, ValueError):
        obj = {"texto": chunk_text}
    digest = row_hash(obj)
    return record_key(obj, key_field, digest), digest
